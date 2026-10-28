import re
import os
import tempfile
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as ExcelImage


def _classify_column(series):
    """Classify a DataFrame column by its data type."""

    if pd.api.types.is_bool_dtype(series):
        return "Boolean"

    if pd.api.types.is_datetime64_any_dtype(series):
        return "Datetime"

    if pd.api.types.is_numeric_dtype(series):
        return "Numeric"

    if (
        pd.api.types.is_object_dtype(series)
        or isinstance(series.dtype, pd.CategoricalDtype)
        or pd.api.types.is_string_dtype(series)
    ):
        return "Categorical"

    return "Other"


def _unique_sheet_name(column, used_names):
    """Create a valid, unique Excel worksheet name."""

    name = re.sub(r'[\[\]:*?/\\]', '_', str(column))
    name = name.strip("'")[:31] or "Column"

    candidate = name
    counter = 1

    while candidate.lower() in used_names:
        suffix = f"_{counter}"
        candidate = name[:31 - len(suffix)] + suffix
        counter += 1

    used_names.add(candidate.lower())

    return candidate


def _create_column_plot(series, column_type, column_name):
    """Create a plot for a column and return its figure."""

    # Skip plots for empty columns.
    values = series.dropna()

    if values.empty:
        return None

    fig, ax = plt.subplots(figsize=(8, 4.5))

    if column_type == "Numeric":

        # Histogram with KDE where sufficient distinct values exist.
        sns.histplot(
            values,
            kde=values.nunique() > 1,
            ax=ax
        )

        ax.set_title(f"Distribution: {column_name}")
        ax.set_xlabel(str(column_name))
        ax.set_ylabel("Frequency")

    elif column_type in ("Categorical", "Boolean"):

        # Display the 20 most frequent values.
        counts = series.value_counts(dropna=False).head(20)

        labels = [
            "<Missing>" if pd.isna(value) else str(value)
            for value in counts.index
        ]

        sns.barplot(
            x=labels,
            y=counts.values,
            ax=ax
        )

        ax.set_title(f"Top Values: {column_name}")
        ax.set_xlabel(str(column_name))
        ax.set_ylabel("Count")
        ax.tick_params(axis="x", rotation=45)

    elif column_type == "Datetime":

        # Display observations by date, when meaningful.
        dates = pd.to_datetime(values, errors="coerce").dropna()

        if dates.empty:
            plt.close(fig)
            return None

        counts = dates.dt.date.value_counts().sort_index()

        ax.plot(
            [str(date) for date in counts.index],
            counts.values,
            marker="o"
        )

        ax.set_title(f"Observations Over Time: {column_name}")
        ax.set_xlabel("Date")
        ax.set_ylabel("Count")
        ax.tick_params(axis="x", rotation=45)

    else:
        plt.close(fig)
        return None

    fig.tight_layout()

    return fig


def _write_section(writer, sheet_name, title, dataframe, startrow):
    """Write a titled section to an Excel worksheet."""

    dataframe.to_excel(
        writer,
        sheet_name=sheet_name,
        index=False,
        startrow=startrow + 1
    )

    worksheet = writer.sheets[sheet_name]

    worksheet.cell(
        row=startrow + 1,
        column=1,
        value=title
    )

    worksheet.cell(
        row=startrow + 1,
        column=1
    ).font = Font(bold=True, size=12)

    return startrow + len(dataframe) + 3


def _format_workbook(writer):
    """Apply basic formatting to all worksheets."""

    for worksheet in writer.sheets.values():

        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False

        for cell in worksheet[1]:

            cell.font = Font(bold=True, color="FFFFFF")

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )

            cell.alignment = Alignment(
                vertical="center",
                horizontal="left",
                wrap_text=True
            )

        # Apply a reasonable column width.
        for column_cells in worksheet.columns:

            column_letter = column_cells[0].column_letter

            max_length = 0

            for cell in column_cells[:100]:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                35
            )


def generate_eda_report(
    df,
    filename="EDA_Report.xlsx",
    plots=True,
    separate_sheets=True
):
    """
    Generate an Excel EDA report from a pandas DataFrame.

    Parameters
    ----------
    df : pandas.DataFrame
        Input dataset.

    filename : str
        Output Excel filename.

    plots : bool
        Include column visualizations when True.

    separate_sheets : bool
        Create an individual worksheet for every column when True.
        Otherwise, write column analyses to a single worksheet.

    Returns
    -------
    str
        Path to the generated Excel report.
    """

    # -------------------------------
    # 1. Validate input
    # -------------------------------

    if not isinstance(df, pd.DataFrame):
        raise TypeError("Input must be a pandas DataFrame.")

    if df.empty:
        raise ValueError("Input DataFrame is empty.")

    if df.shape[1] == 0:
        raise ValueError("DataFrame has no columns.")

    if not df.columns.is_unique:
        raise ValueError(
            "DataFrame contains duplicate column names. "
            "Rename duplicate columns before generating the report."
        )

    output_dir = os.path.dirname(os.path.abspath(filename))
    os.makedirs(output_dir, exist_ok=True)

    # Used to manage temporary plot images.
    with tempfile.TemporaryDirectory() as temp_dir:

        # -------------------------------
        # 2. Dataset-level summary
        # -------------------------------

        summary = pd.DataFrame({
            "Column": df.columns.astype(str),
            "Data Type": [str(dtype) for dtype in df.dtypes],
            "Classification": [
                _classify_column(df[col])
                for col in df.columns
            ],
            "Missing Values": df.isna().sum().values,
            "Missing Percentage": (
                df.isna().mean().values * 100
            ).round(2),
            "Unique Values": [
                df[col].nunique()
                for col in df.columns
            ]
        })

        dataset_info = pd.DataFrame({
            "Metric": [
                "Total Rows",
                "Total Columns",
                "Duplicate Rows",
                "Total Missing Values"
            ],
            "Value": [
                len(df),
                len(df.columns),
                int(df.duplicated().sum()),
                int(df.isna().sum().sum())
            ]
        })

        # -------------------------------
        # 3. Create Excel workbook
        # -------------------------------

        used_names = {"summary"}

        with pd.ExcelWriter(
            filename,
            engine="openpyxl"
        ) as writer:

            # Overall dataset information.
            dataset_info.to_excel(
                writer,
                sheet_name="Summary",
                index=False,
                startrow=0
            )

            _write_section(
                writer,
                "Summary",
                "Column Overview",
                summary,
                startrow=len(dataset_info) + 3
            )

            # -------------------------------
            # 4. Analyze each column
            # -------------------------------

            combined_sheet = "Column Analysis"

            if not separate_sheets:
                used_names.add(combined_sheet.lower())

            combined_startrow = 0

            if not separate_sheets:
                # Create an initial worksheet for all columns.
                pd.DataFrame({
                    "Report": ["Column Analysis"]
                }).to_excel(
                    writer,
                    sheet_name=combined_sheet,
                    index=False
                )

            for column in df.columns:

                series = df[column]
                column_type = _classify_column(series)

                if separate_sheets:
                    sheet_name = _unique_sheet_name(
                        column,
                        used_names
                    )
                else:
                    sheet_name = combined_sheet

                startrow = (
                    0 if separate_sheets
                    else combined_startrow
                )

                # Column metadata.
                metadata = pd.DataFrame({
                    "Metric": [
                        "Column Name",
                        "Data Type",
                        "Classification",
                        "Total Rows",
                        "Missing Values",
                        "Missing Percentage",
                        "Unique Values"
                    ],
                    "Value": [
                        str(column),
                        str(series.dtype),
                        column_type,
                        len(series),
                        int(series.isna().sum()),
                        round(series.isna().mean() * 100, 2),
                        int(series.nunique())
                    ]
                })

                startrow = _write_section(
                    writer,
                    sheet_name,
                    f"Column Profile: {column}",
                    metadata,
                    startrow
                )

                # Descriptive statistics.
                if column_type == "Numeric":

                    stats = series.describe().reset_index()
                    stats.columns = ["Statistic", "Value"]
                    outlier_count, outlier_percentage = self_outlier_count(series)
     
                    # Additional numeric information.
                    extra_stats = pd.DataFrame({
                        "Statistic": [
                            "Median",
                            "Variance",
                            "Skewness",
                            "Kurtosis",
                            "Zero Count",
                            "Negative Count",
                            "Outlier Count (IQR)",
                            "Outlier Percentage (%)"

                        ],
                        "Value": [
                            series.median(),
                            series.var(),
                            series.skew(),
                            series.kurt(),
                            int((series == 0).sum()),
                            int((series < 0).sum()),
                            outlier_count,
                            outlier_percentage                        ]
                    })

                    stats = pd.concat(
                        [stats, extra_stats],
                        ignore_index=True
                    )

                elif column_type == "Datetime":

                    dates = pd.to_datetime(
                        series,
                        errors="coerce"
                    )

                    stats = pd.DataFrame({
                        "Statistic": [
                            "Count",
                            "Minimum Date",
                            "Maximum Date",
                            "Unique Dates"
                        ],
                        "Value": [
                            dates.notna().sum(),
                            dates.min(),
                            dates.max(),
                            dates.nunique()
                        ]
                    })

                else:

                    stats = pd.DataFrame({
                        "Statistic": [
                            "Non-Missing Count",
                            "Unique Values"
                        ],
                        "Value": [
                            series.notna().sum(),
                            series.nunique()
                        ]
                    })

                startrow = _write_section(
                    writer,
                    sheet_name,
                    "Descriptive Statistics",
                    stats,
                    startrow
                )

                # Value counts, including missing values.
                counts = (
                    series.value_counts(dropna=False)
                    .head(20)
                    .rename_axis("Value")
                    .reset_index(name="Count")
                )

                counts["Percentage"] = (
                    counts["Count"] / len(series) * 100
                ).round(2)

                startrow = _write_section(
                    writer,
                    sheet_name,
                    "Value Counts",
                    counts,
                    startrow
                )

                # -------------------------------
                # 5. Generate and embed plots
                # -------------------------------

                if plots:

                    fig = _create_column_plot(
                        series,
                        column_type,
                        column
                    )

                    if fig is not None:

                        image_path = os.path.join(
                            temp_dir,
                            f"plot_{len(os.listdir(temp_dir))}.png"
                        )

                        fig.savefig(
                            image_path,
                            dpi=120,
                            bbox_inches="tight"
                        )

                        plt.close(fig)

                        worksheet = writer.sheets[sheet_name]

                        image = ExcelImage(image_path)

                        image.width = 600
                        image.height = 340

                        image_cell = f"A{startrow + 2}"

                        worksheet.add_image(
                            image,
                            image_cell
                        )

                        if not separate_sheets:
                            startrow += 22

                if not separate_sheets:
                    combined_startrow = startrow + 3

            # -------------------------------
            # 6. Format workbook
            # -------------------------------

            _format_workbook(writer)

    return os.path.abspath(filename)


def self_outlier_count(series):
    """Count IQR outliers, excluding binary numeric columns."""

    series = series.dropna()

    if series.empty or series.nunique() <= 2:
        return 0, 0.0

    Q1 = series.quantile(0.25)
    Q3 = series.quantile(0.75)
    IQR = Q3 - Q1

    lower_limit = Q1 - 1.5 * IQR
    upper_limit = Q3 + 1.5 * IQR

    outlier_count = (
        (series < lower_limit) | (series > upper_limit)
    ).sum()

    outlier_percentage = (outlier_count / len(series)) * 100

    return int(outlier_count), round(outlier_percentage, 2)
